"""Gabarito de janeiro: por dupla (valores sorteados) e do caso-base (para conferência do professor)."""
from __future__ import annotations

import io
from urllib.parse import quote

from fastapi import APIRouter, Depends, Request
from fastapi.responses import Response
from sqlalchemy.orm import Session

from .. import plano_contas as pc
from ..db import get_db
from ..gabarito import gerar_gabarito
from ..geracao import gerar
from ..main import render, voltar
from ..models_caso import Caso
from ..models_dupla import Dupla

router = APIRouter()
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _eventos(caso: Caso | None) -> dict:
    out = {"TES": "Tesouraria diária (resgates e conta garantida)", "AJ": "Ajustes de fim de mês (apropriações e juros)",
           "J40": "Apuração dos tributos de janeiro"}
    if caso:
        for e in caso.eventos:
            if e.data and e.data.month == 1:
                out.setdefault(e.codigo, e.descricao)
                out[e.codigo] = e.descricao
    return out


def _grupos(g: dict) -> list[dict]:
    """Lançamentos agrupados por evento, na ordem em que ocorrem."""
    ordem, por_ev = [], {}
    for l in g["lancamentos"]:
        if l["evento"] not in por_ev:
            por_ev[l["evento"]] = {"evento": l["evento"], "data": l["data"], "lancs": []}
            ordem.append(l["evento"])
        por_ev[l["evento"]]["lancs"].append(l)
    return [por_ev[e] for e in ordem]


def _balancete(g: dict) -> list[dict]:
    """Analíticas com linhas de grupo (sintéticas) para leitura."""
    linhas, vistos = [], set()
    soma = {}
    for x in g["balancete"]:
        for s in pc.sinteticas_de(x["conta"]):
            t = soma.setdefault(s, {"inicial": 0.0, "debitos": 0.0, "creditos": 0.0, "final": 0.0})
            for k in t:
                t[k] += x[k]
    for x in g["balancete"]:
        for s in pc.sinteticas_de(x["conta"]):
            if s not in vistos and s.count(".") <= 2:
                vistos.add(s)
                linhas.append({"conta": s, "nome": pc.SINTETICAS[s], "sintetica": True, "nivel": s.count("."),
                               **{k: round(v, 2) for k, v in soma[s].items()}})
        linhas.append({**x, "sintetica": False, "nivel": 4})
    return linhas


STATUS = {"PENDENTE": "Pendente", "VALIDADA": "Validada", "AJUSTAR": "Ajustar"}


def _contexto(g: dict, caso: Caso | None, editar: bool = False) -> dict:
    val = ((caso.dados or {}).get("premissas_validacao") or {}) if caso else {}
    pend = sum(1 for k in g["premissas"] if (val.get(k) or {}).get("status", "PENDENTE") != "VALIDADA")
    return {"g": g, "grupos": _grupos(g), "linhas_bal": _balancete(g), "eventos": _eventos(caso),
            "validacao": val, "status_opcoes": STATUS, "pendentes": pend, "editar_premissas": editar, "caso_ref": caso}


@router.get("/duplas/{id}/gabarito")
def gabarito_dupla(request: Request, id: int, db: Session = Depends(get_db)):
    d = db.get(Dupla, id)
    if not d or not d.empresa:
        return voltar("/duplas", erro="Gere a empresa da dupla antes do gabarito.")
    g = gerar_gabarito(d.empresa.valores)
    return render(request, "gabarito/ver.html", d=d, titulo=f"{d.empresa.valores['empresa'].get('nome', 'Empresa')} · {d.rotulo}",
                  xlsx=f"/duplas/{id}/gabarito.xlsx", voltar_url=f"/duplas/{id}/empresa", menu="duplas", **_contexto(g, d.empresa.caso))


@router.get("/casos/{id}/gabarito")
def gabarito_caso(request: Request, id: int, db: Session = Depends(get_db)):
    from .casos import ABAS, ESCOPOS, NATUREZAS
    caso = db.get(Caso, id)
    if not caso:
        return voltar("/casos", erro="Caso não encontrado.")
    g = gerar_gabarito(gerar(caso, 0, "Caso-base", None, base=True))
    return render(request, "gabarito/caso.html", caso=caso, aba="gabarito", abas=ABAS, escopos=ESCOPOS, naturezas=NATUREZAS,
                  xlsx=f"/casos/{id}/gabarito.xlsx", menu="casos", **_contexto(g, caso, editar=True))


@router.post("/casos/{id}/gabarito/premissas")
async def validar_premissas(request: Request, id: int, db: Session = Depends(get_db)):
    """Validação das premissas do gabarito pelo professor, gravada no próprio caso."""
    from datetime import datetime
    caso = db.get(Caso, id)
    if not caso:
        return voltar("/casos", erro="Caso não encontrado.")
    form = await request.form()
    dados = dict(caso.dados or {})
    val = dict(dados.get("premissas_validacao") or {})
    for k in [x for x in form.keys() if x.startswith("st_")]:
        cod = k[3:]
        st = form.get(k) if form.get(k) in STATUS else "PENDENTE"
        com = (form.get(f"com_{cod}") or "").strip()
        ant = val.get(cod) or {}
        if ant.get("status") != st or ant.get("comentario", "") != com:
            val[cod] = {"status": st, "comentario": com, "em": datetime.now().strftime("%d/%m/%Y %H:%M")}
    dados["premissas_validacao"] = val
    caso.dados = dados
    db.commit()
    return voltar(f"/casos/{id}/gabarito#premissas", "Validação das premissas salva.")


def _xlsx(g: dict, titulo: str, eventos: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    cab = Font(bold=True, color="FFFFFF")
    fundo = PatternFill("solid", fgColor="2B59D1")
    num = '#,##0.00;[Red]-#,##0.00'

    def cabecalho(ws, cols, larg):
        ws.append(cols)
        for i, c in enumerate(ws[ws.max_row], 1):
            c.font, c.fill = cab, fundo
            ws.column_dimensions[c.column_letter].width = larg[i - 1]
        ws.freeze_panes = ws.cell(row=ws.max_row + 1, column=1)

    ws = wb.active
    ws.title = "Diário"
    ws.append([titulo])
    ws["A1"].font = Font(bold=True, size=13)
    cabecalho(ws, ["Nº", "Data", "Evento", "Histórico", "Conta", "Nome da conta", "Débito", "Crédito", "Premissas"],
              [6, 11, 8, 46, 22, 46, 15, 15, 14])
    for l in g["lancamentos"]:
        for p in l["debitos"]:
            ws.append([l["n"], l["data"][8:10] + "/" + l["data"][5:7], l["evento"], l["historico"], p["conta"], p["nome"], p["valor"], None,
                       ", ".join(l["premissas"])])
        for p in l["creditos"]:
            ws.append([l["n"], l["data"][8:10] + "/" + l["data"][5:7], l["evento"], l["historico"], p["conta"], p["nome"], None, p["valor"],
                       ", ".join(l["premissas"])])
    for row in ws.iter_rows(min_row=3, min_col=7, max_col=8):
        for c in row:
            c.number_format = num

    ws = wb.create_sheet("Balancete")
    cabecalho(ws, ["Conta", "Nome", "Saldo inicial", "Débitos", "Créditos", "Saldo final", "D/C"], [22, 52, 16, 16, 16, 16, 5])
    for x in g["balancete"]:
        ws.append([x["conta"], x["nome"], abs(x["inicial"]), x["debitos"], x["creditos"], abs(x["final"]),
                   "D" if x["final"] >= 0 else "C"])
    t = g["totais"]
    ws.append([])
    ws.append(["", "Total devedor / credor", None, None, None, t["devedor"], ""])
    ws.append(["", "", None, None, None, t["credor"], ""])
    ws.append(["", "Resultado de janeiro (antes do encerramento)", None, None, None, t["resultado"], ""])
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=6):
        for c in row:
            c.number_format = num

    ws = wb.create_sheet("Memória de cálculo")
    cabecalho(ws, ["Evento", "Descrição do evento", "Item", "Valor", "Observação"], [8, 40, 90, 16, 26])
    for ev, itens in g["memoria"].items():
        for m in itens:
            ws.append([ev, eventos.get(ev, ""), m["descricao"], m["valor"], m["obs"]])
    for row in ws.iter_rows(min_row=2, min_col=4, max_col=4):
        for c in row:
            c.number_format = num

    ws = wb.create_sheet("Premissas")
    cabecalho(ws, ["Código", "Premissa (validar)", "Validada?", "Comentário do professor"], [8, 120, 12, 50])
    for k, v in g["premissas"].items():
        ws.append([k, v, "", ""])
        ws.cell(row=ws.max_row, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@router.get("/duplas/{id}/gabarito.xlsx")
def gabarito_dupla_xlsx(id: int, db: Session = Depends(get_db)):
    d = db.get(Dupla, id)
    if not d or not d.empresa:
        return voltar("/duplas", erro="Gere a empresa da dupla antes do gabarito.")
    g = gerar_gabarito(d.empresa.valores)
    nome = f"Gabarito janeiro - {d.turma.disciplina.codigo} {d.turma.codigo} - Dupla {d.numero:02d}.xlsx"
    corpo = _xlsx(g, f"Gabarito de janeiro/2026 · {d.empresa.valores['empresa'].get('nome', '')} · {d.rotulo}", _eventos(d.empresa.caso))
    return Response(corpo, media_type=XLSX, headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(nome)}"})


@router.get("/casos/{id}/gabarito.xlsx")
def gabarito_caso_xlsx(id: int, db: Session = Depends(get_db)):
    caso = db.get(Caso, id)
    if not caso:
        return voltar("/casos", erro="Caso não encontrado.")
    g = gerar_gabarito(gerar(caso, 0, "Caso-base", None, base=True))
    nome = f"Gabarito janeiro - caso-base {caso.id}.xlsx"
    corpo = _xlsx(g, f"Gabarito de janeiro/2026 · {caso.nome} (valores do caso-base)", _eventos(caso))
    return Response(corpo, media_type=XLSX, headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(nome)}"})
